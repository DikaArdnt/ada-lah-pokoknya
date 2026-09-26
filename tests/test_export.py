from __future__ import annotations

import re
import zipfile
from dataclasses import replace
from pathlib import Path

import docx as docx_lib
import pytest

from ocr_reconstructor.config import WatermarkConfig
from ocr_reconstructor.core.exceptions import ExportError
from ocr_reconstructor.documents.docx import export_docx
from ocr_reconstructor.documents.pdf import export_pdf

SAMPLE_MD = """# Annual Report

Intro paragraph with **bold** and *italic* and `code`.

## Details

- first bullet
- second bullet

1. numbered one
2. numbered two

> quoted line

---

Final paragraph.
"""


def test_docx_export_contains_selectable_text(tmp_path: Path):
    output = tmp_path / "document.docx"
    export_docx(SAMPLE_MD, output)
    assert output.is_file()

    document = docx_lib.Document(str(output))
    paragraphs = [p.text for p in document.paragraphs]
    assert "Annual Report" in paragraphs
    assert any("Intro paragraph with" in text for text in paragraphs)
    assert any("first bullet" in text for text in paragraphs)
    assert any("numbered one" in text for text in paragraphs)
    assert any("quoted line" in text for text in paragraphs)

    # Inline formatting became runs.
    intro = next(p for p in document.paragraphs if "Intro paragraph" in p.text)
    bold_runs = [r.text for r in intro.runs if r.bold]
    assert bold_runs == ["bold"]


def test_docx_multiple_choice_options_on_separate_lines(tmp_path: Path):
    md = "1. Soal pertama\n\n   A. opsi satu\n   B. opsi dua\n   C. opsi tiga\n"
    output = tmp_path / "choices.docx"
    export_docx(md, output)
    document = docx_lib.Document(str(output))
    paragraphs = [p.text for p in document.paragraphs]
    assert "A. opsi satu" in paragraphs
    assert "B. opsi dua" in paragraphs
    assert "C. opsi tiga" in paragraphs
    assert not any("opsi satu opsi dua" in text for text in paragraphs)


def test_docx_heading_levels(tmp_path: Path):
    output = tmp_path / "document.docx"
    export_docx(SAMPLE_MD, output)
    document = docx_lib.Document(str(output))
    heading = next(p for p in document.paragraphs if p.text == "Annual Report")
    assert heading.style.name == "Heading 1"
    section = next(p for p in document.paragraphs if p.text == "Details")
    assert section.style.name == "Heading 2"


def test_docx_watermark_writes_header_shape(tmp_path: Path):
    output = tmp_path / "document.docx"
    watermark = WatermarkConfig(enabled=True, text="DRAFT")
    export_docx(SAMPLE_MD, output, watermark)

    with zipfile.ZipFile(output) as archive:
        header_files = [n for n in archive.namelist() if n.startswith("word/header")]
        assert header_files, "no header part written"
        content = archive.read(header_files[0]).decode("utf-8")
    assert "textpath" in content
    assert "DRAFT" in content


def test_pdf_export_produces_valid_pdf(tmp_path: Path):
    output = tmp_path / "document.pdf"
    export_pdf(SAMPLE_MD, output)
    raw = output.read_bytes()
    assert raw.startswith(b"%PDF-")
    assert b"/Type /Page" in raw  # uncompressed streams are inspectable
    assert b"Annual Report" in raw  # selectable text content


def test_pdf_watermark_draws_text(tmp_path: Path):
    output = tmp_path / "document.pdf"
    watermark = WatermarkConfig(enabled=True, text="CONFIDENTIAL", opacity=0.3)
    export_pdf(SAMPLE_MD, output, watermark)
    raw = output.read_bytes()
    assert b"CONFIDENTIAL" in raw


def test_pdf_tiled_watermark(tmp_path: Path):
    output = tmp_path / "document.pdf"
    watermark = WatermarkConfig(enabled=True, text="TILE", position="tile")
    export_pdf(SAMPLE_MD, output, watermark)
    assert b"TILE" in output.read_bytes()


def test_export_without_watermark_has_no_stamp(tmp_path: Path):
    output = tmp_path / "document.pdf"
    export_pdf(SAMPLE_MD, output, WatermarkConfig(enabled=False))
    assert b"DRAFT" not in output.read_bytes()


# ---------------------------------------------------------------------------
# Precise watermark positioning
# ---------------------------------------------------------------------------

ALL_POSITIONS = [
    "center",
    "horizontal",
    "vertical",
    "diagonal",
    "tile",
    "top-left",
    "top",
    "top-right",
    "left",
    "right",
    "bottom-left",
    "bottom",
    "bottom-right",
]


def test_rotation_only_applies_to_diagonal():
    from ocr_reconstructor.documents.watermark import _rotation_for

    base = WatermarkConfig(enabled=True, text="DRAFT", rotation=45)
    assert _rotation_for(base) == 0  # center stays horizontal
    assert _rotation_for(replace(base, position="diagonal")) == 45
    assert _rotation_for(replace(base, position="vertical")) == 90
    assert _rotation_for(replace(base, position="tile")) == 0
    assert _rotation_for(replace(base, position="top-left", rotation=45)) == 0


def test_anchor_factors_cover_the_page():
    from ocr_reconstructor.documents.watermark import _anchor

    assert _anchor(WatermarkConfig(position="center")) == (0.5, 0.5)
    assert _anchor(WatermarkConfig(position="diagonal")) == (0.5, 0.5)
    assert _anchor(WatermarkConfig(position="vertical")) == (0.5, 0.5)
    assert _anchor(WatermarkConfig(position="horizontal")) == (0.5, 0.5)
    assert _anchor(WatermarkConfig(position="top-left")) == (0.0, 1.0)
    assert _anchor(WatermarkConfig(position="bottom-right")) == (1.0, 0.0)
    assert _anchor(WatermarkConfig(position="left")) == (0.0, 0.5)
    assert _anchor(WatermarkConfig(position="right")) == (1.0, 0.5)
    assert _anchor(WatermarkConfig(position="top")) == (0.5, 1.0)
    assert _anchor(WatermarkConfig(position="bottom")) == (0.5, 0.0)


@pytest.mark.parametrize("position", ALL_POSITIONS)
def test_all_positions_export_valid_docx_and_pdf(tmp_path: Path, position: str):
    watermark = WatermarkConfig(enabled=True, text="DRAFT", position=position)
    docx_out = tmp_path / f"{position}.docx"
    pdf_out = tmp_path / f"{position}.pdf"
    export_docx(SAMPLE_MD, docx_out, watermark)
    export_pdf(SAMPLE_MD, pdf_out, watermark)
    assert docx_out.is_file() and pdf_out.is_file()
    assert b"DRAFT" in pdf_out.read_bytes()


@pytest.mark.parametrize("position", ALL_POSITIONS)
def test_docx_watermark_style_matches_position(tmp_path: Path, position: str):
    output = tmp_path / f"{position}.docx"
    export_docx(SAMPLE_MD, output, WatermarkConfig(enabled=True, text="DRAFT", position=position))
    with zipfile.ZipFile(output) as archive:
        header_files = [n for n in archive.namelist() if n.startswith("word/header")]
        content = archive.read(header_files[0]).decode("utf-8")
    style = re.search(r'<v:shape [^>]*style="([^"]+)"', content).group(1)
    # Anchoring is relative to the page, not the text margins.
    assert "mso-position-horizontal-relative:page" in style
    assert "mso-position-vertical-relative:page" in style
    if position == "center" or position == "horizontal":
        assert "mso-position-horizontal:center" in style
        assert "mso-position-vertical:center" in style
        assert "rotation:0" in style
    elif position == "diagonal":
        assert "mso-position-horizontal:center" in style
        assert "rotation:45" in style
    elif position == "vertical":
        assert "rotation:90" in style
    elif position == "top-left":
        assert "mso-position-horizontal:left" in style
        assert "mso-position-vertical:top" in style
    elif position == "bottom-right":
        assert "mso-position-horizontal:right" in style
        assert "mso-position-vertical:bottom" in style
    # The shape size follows the text so font_size is honoured.
    assert re.search(r"width:[\d.]+pt;height:[\d.]+pt", style)


def test_docx_custom_diagonal_rotation(tmp_path: Path):
    output = tmp_path / "document.docx"
    export_docx(
        SAMPLE_MD,
        output,
        WatermarkConfig(enabled=True, text="MIRING", position="diagonal", rotation=30),
    )
    with zipfile.ZipFile(output) as archive:
        header_files = [n for n in archive.namelist() if n.startswith("word/header")]
        content = archive.read(header_files[0]).decode("utf-8")
    assert "rotation:30" in content


def test_docx_tile_writes_a_stamp_grid(tmp_path: Path):
    output = tmp_path / "document.docx"
    export_docx(SAMPLE_MD, output, WatermarkConfig(enabled=True, text="TILE", position="tile"))
    with zipfile.ZipFile(output) as archive:
        header_files = [n for n in archive.namelist() if n.startswith("word/header")]
        content = archive.read(header_files[0]).decode("utf-8")
    shapes = re.findall(r"<v:shape ", content)
    assert len(shapes) > 4, "tile should repeat several stamps"
    # The shapetype is declared once, later shapes only reference it.
    assert len(re.findall(r"<v:shapetype ", content)) == 1


def test_pdf_center_watermark_is_at_page_center(tmp_path: Path):
    output = tmp_path / "document.pdf"
    export_pdf(SAMPLE_MD, output, WatermarkConfig(enabled=True, text="DRAFT", position="center"))
    raw = output.read_bytes().decode("latin-1")
    match = re.search(r"1 0 0 1 ([\d.]+) ([\d.]+) cm\nBT 1 0 0 1", raw)
    assert match
    x, y = float(match.group(1)), float(match.group(2))
    assert x == pytest.approx(595.2756 / 2, abs=0.5)  # A4 width / 2
    assert y == pytest.approx(841.8898 / 2, abs=0.5)  # A4 height / 2


def test_pdf_diagonal_rotates_around_page_center(tmp_path: Path):
    output = tmp_path / "document.pdf"
    export_pdf(SAMPLE_MD, output, WatermarkConfig(enabled=True, text="DRAFT", position="diagonal"))
    raw = output.read_bytes().decode("latin-1")
    # A rotation matrix (cos, sin, -sin, cos) for -45deg, centered on A4.
    match = re.search(
        r"\.707\d* (-?0?\.707\d*) \.707\d* \.707\d* (297\.\d+) (420\.\d+)", raw
    )
    assert match


def test_pdf_top_left_uses_edge_inset(tmp_path: Path):
    output = tmp_path / "document.pdf"
    export_pdf(
        SAMPLE_MD,
        output,
        WatermarkConfig(enabled=True, text="DRAFT", position="top-left"),
    )
    raw = output.read_bytes().decode("latin-1")
    match = re.search(r"1 0 0 1 ([\d.]+) ([\d.]+) cm\nBT 1 0 0 1 0 -\d+", raw)
    assert match
    x, y = float(match.group(1)), float(match.group(2))
    inset = 1.5 * 28.3465  # 1.5 cm in points
    assert x == pytest.approx(inset, abs=0.5)
    assert y == pytest.approx(841.8898 - inset, abs=0.5)


def test_pdf_vertical_rotates_90_degrees(tmp_path: Path):
    output = tmp_path / "document.pdf"
    export_pdf(
        SAMPLE_MD,
        output,
        WatermarkConfig(enabled=True, text="DRAFT", position="vertical"),
    )
    raw = output.read_bytes().decode("latin-1")
    assert re.search(r"0 -1 1 0 (297\.\d+) (420\.\d+)", raw)
