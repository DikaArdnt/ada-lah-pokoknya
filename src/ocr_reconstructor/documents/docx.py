from __future__ import annotations

from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ocr_reconstructor.config import WatermarkConfig
from ocr_reconstructor.core.exceptions import ExportError
from ocr_reconstructor.documents.markdown import Block, parse_inline, parse_markdown
from ocr_reconstructor.documents.watermark import apply_docx_watermark
from ocr_reconstructor.utils.filesystem import ensure_dir

_CODE_FONT = "Consolas"


def _add_runs(paragraph, text: str) -> None:
    for span in parse_inline(text):
        run = paragraph.add_run(span.text)
        run.bold = span.bold or None
        run.italic = span.italic or None
        if span.code:
            run.font.name = _CODE_FONT


def _add_bottom_border(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "auto")
    borders.append(bottom)
    p_pr.append(borders)


def _render_block(document, block: Block) -> None:
    if block.kind == "heading":
        level = min(block.level, 6)
        document.add_heading("", level=level)
        heading_paragraph = document.paragraphs[-1]
        _add_runs(heading_paragraph, block.text)
    elif block.kind == "paragraph":
        paragraph = document.add_paragraph()
        _add_runs(paragraph, block.text)
    elif block.kind == "list_item":
        if block.ordered:
            paragraph = document.add_paragraph()
            _add_runs(paragraph, f"{block.number}{block.delimiter} {block.text}")
        else:
            paragraph = document.add_paragraph(style="List Bullet")
            _add_runs(paragraph, block.text)
    elif block.kind == "quote":
        paragraph = document.add_paragraph(style="Intense Quote")
        _add_runs(paragraph, block.text)
    elif block.kind == "rule":
        paragraph = document.add_paragraph()
        _add_bottom_border(paragraph)


def export_docx(
    markdown_text: str,
    output_path: Path,
    watermark: WatermarkConfig | None = None,
    title: str | None = None,
) -> Path:
    try:
        document = docx.Document()
        for block in parse_markdown(markdown_text):
            _render_block(document, block)
        document.core_properties.title = title or output_path.stem
        if watermark is not None and watermark.enabled:
            apply_docx_watermark(document, watermark)
        ensure_dir(output_path.parent)
        document.save(str(output_path))
    except ExportError:
        raise
    except Exception as exc:
        raise ExportError(
            f"DOCX export failed: {exc}",
            "Check that the output directory is writable and the document "
            "Markdown is well-formed.",
        ) from exc
    return output_path
